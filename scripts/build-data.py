from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = ROOT / "data.xlsx"
DATABASE_FILE = ROOT / "whumath_database.xlsx"
JSON_FILE = ROOT / "data" / "teamData.json"
JS_FILE = ROOT / "data" / "teamData.js"

TEAM_CN = "数院"
TEAM_EN = "WHUMATH"
TEAM_FULL_CN = "武大数院"

STARTERS_BY_COMPETITION = {
    "新生杯": 8,
    "振兴杯": 11,
    "五人制": 5,
}

FORMAT_BY_COMPETITION = {
    "新生杯": "8-a-side",
    "振兴杯": "11-a-side",
    "五人制": "5-a-side",
}

ACHIEVEMENTS = [
    {
        "year": 2024,
        "competition": "五人制",
        "title": "2024年五人制冠军",
        "titleEn": "2024 Five-a-side Champion",
        "image": "assets/images/champion.jpg",
    },
    {
        "year": 2026,
        "competition": "五人制",
        "title": "2026年五人制亚军",
        "titleEn": "2026 Five-a-side Runner-up",
        "image": "",
    },
]

SEASON_CAPTAINS = {
    "2023-2024": "胡然",
    "2024-2025": "廖梓瑞",
    "2025-2026": "李昊哲",
}

JERSEY_NUMBERS = {
    "2023-2024": {"李昊哲": "10"},
    "2024-2025": {"李昊哲": "10"},
    "2025-2026": {"李昊哲": "10"},
}

PLAYER_POSITIONS = {
    "李昊哲": "前腰",
}

VENUE_NAMES = {
    "桂操": "桂园操场",
    "信操": "信息学部操场",
    "工操": "工学部操场",
    "奥场": "九一二操场",
}

NEWS_ITEMS = [
    {
        "season": "2025-2026",
        "type": "新闻链接",
        "typeEn": "News",
        "title": "球队新闻链接",
        "titleEn": "Team news links",
        "date": "",
        "url": "",
        "cover": "",
        "description": "待补充",
        "descriptionEn": "To be added",
    },
    {
        "season": "2025-2026",
        "type": "采访录像",
        "typeEn": "Interview",
        "title": "采访录像",
        "titleEn": "Interview videos",
        "date": "",
        "url": "",
        "cover": "",
        "description": "待补充",
        "descriptionEn": "To be added",
    },
]

TRANSFERS = [
    {
        "season": "2025-2026",
        "newcomers": [
            {"name": "贺宇洲", "date": "2025.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "王嘉宁", "date": "2025.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "应锦程", "date": "2025.09", "from": "吉林大学", "to": TEAM_FULL_CN},
            {"name": "孔维轩", "date": "2024.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "姚宇辰", "date": "2024.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "刘岳然", "date": "2024.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "张硕铠", "date": "2024.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "姚健", "date": "2024.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "方烁俨", "date": "2025.02", "from": "武大经管", "to": TEAM_FULL_CN},
            {"name": "李昊哲", "date": "2024.02", "from": "武大物院", "to": TEAM_FULL_CN},
            {"name": "许景粟", "date": "2023.09", "from": "武大化院", "to": TEAM_FULL_CN},
            {"name": "刘恺源", "date": "2023.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "韩志豪", "date": "2023.09", "from": "青训", "to": TEAM_FULL_CN},
            {"name": "邝钧皓", "date": "2023.09", "from": "青训", "to": TEAM_FULL_CN},
        ],
        "departures": [
            {"name": "邝钧皓", "date": "2025.09", "from": TEAM_FULL_CN, "to": "退役"},
            {"name": "周宇杰", "date": "2024.06", "from": TEAM_FULL_CN, "to": "澳门大学"},
            {"name": "胡然", "date": "2024.06", "from": TEAM_FULL_CN, "to": "哥伦比亚大学"},
            {"name": "石德铭", "date": "2024.06", "from": TEAM_FULL_CN, "to": "南洋理工大学"},
            {"name": "格桑次仁", "date": "2024.06", "from": TEAM_FULL_CN, "to": "退役"},
            {"name": "遵追桑布", "date": "2024.06", "from": TEAM_FULL_CN, "to": "退役"},
            {"name": "蒋思杨", "date": "2026.06", "from": TEAM_FULL_CN, "to": "退役"},
            {"name": "廖梓瑞", "date": "2026.06", "from": TEAM_FULL_CN, "to": "新加坡国立大学"},
            {"name": "徐锋庆", "date": "2026.06", "from": TEAM_FULL_CN, "to": "退役"},
            {"name": "张梦洋", "date": "2026.06", "from": TEAM_FULL_CN, "to": "退役"},
        ],
    }
]


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y.%m.%d")
    return str(value).strip()


def slugify(*parts: str) -> str:
    text = "-".join(part for part in parts if part)
    text = re.sub(r"\s+", "-", text.strip())
    text = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", text)
    return text.strip("-").lower()


def normalize_name(name: str) -> str:
    name = re.sub(r"\s+", "", name.strip())
    return name


def full_venue_name(venue: str) -> str:
    return VENUE_NAMES.get(venue, venue)


def player_number(season: str, name: str) -> str:
    return JERSEY_NUMBERS.get(season, {}).get(name, "")


def player_position(name: str) -> str:
    return PLAYER_POSITIONS.get(name, "")


def parse_score(score_text: str) -> dict:
    score_text = clean_text(score_text)
    parsed = {
        "raw": score_text,
        "goalsFor": None,
        "goalsAgainst": None,
        "penaltyFor": None,
        "penaltyAgainst": None,
        "result": "unknown",
        "known": False,
    }
    if not score_text:
        return parsed

    match = re.search(
        r"(\d+)\s*[:：]\s*(\d+)(?:\s*[（(]\s*(\d+)\s*[:：]\s*(\d+)\s*[）)])?",
        score_text,
    )
    if not match:
        return parsed

    goals_for = int(match.group(1))
    goals_against = int(match.group(2))
    penalty_for = int(match.group(3)) if match.group(3) is not None else None
    penalty_against = int(match.group(4)) if match.group(4) is not None else None

    result = "draw"
    if penalty_for is not None and penalty_against is not None and goals_for == goals_against:
        if penalty_for > penalty_against:
            result = "win"
        elif penalty_for < penalty_against:
            result = "loss"
    elif goals_for > goals_against:
        result = "win"
    elif goals_for < goals_against:
        result = "loss"

    parsed.update(
        {
            "goalsFor": goals_for,
            "goalsAgainst": goals_against,
            "penaltyFor": penalty_for,
            "penaltyAgainst": penalty_against,
            "result": result,
            "known": True,
        }
    )
    return parsed


def parse_people_list(roster_text: str, starters_count: int) -> list[dict]:
    if not roster_text:
        return []

    raw_people = re.split(r"[、,，\n]+", roster_text)
    people = []
    for raw in raw_people:
        item = clean_text(raw)
        if not item:
            continue

        captain = bool(re.search(r"\(\s*C\s*\)|（\s*C\s*）", item, flags=re.IGNORECASE))
        goalkeeper = bool(re.search(r"\(\s*GK\s*\)|（\s*GK\s*）", item, flags=re.IGNORECASE))
        penalty_marker = bool(re.search(r"\(\s*P\s*\)|（\s*P\s*）", item, flags=re.IGNORECASE))
        name = re.sub(r"[（(]\s*(C|GK|P)\s*[）)]", "", item, flags=re.IGNORECASE)
        name = normalize_name(name)
        if not name:
            continue

        index = len(people)
        people.append(
            {
                "name": name,
                "captain": captain,
                "goalkeeper": goalkeeper,
                "penaltyMarker": penalty_marker,
                "starter": index < starters_count,
            }
        )
    return people


def parse_person_item(raw: str, starter: bool) -> dict | None:
    item = clean_text(raw)
    if not item:
        return None
    captain = bool(re.search(r"\(\s*C\s*\)|（\s*C\s*）", item, flags=re.IGNORECASE))
    goalkeeper = bool(re.search(r"\(\s*GK\s*\)|（\s*GK\s*）", item, flags=re.IGNORECASE))
    penalty_marker = bool(re.search(r"\(\s*P\s*\)|（\s*P\s*）", item, flags=re.IGNORECASE))
    name = re.sub(r"[（(]\s*(C|GK|P)\s*[）)]", "", item, flags=re.IGNORECASE)
    name = normalize_name(name)
    if not name:
        return None
    return {
        "name": name,
        "captain": captain,
        "goalkeeper": goalkeeper,
        "penaltyMarker": penalty_marker,
        "starter": starter,
    }


def parse_roster_columns(starting_text: str, substitutes_text: str) -> list[dict]:
    people = []
    for raw in re.split(r"[、,，\n]+", starting_text or ""):
        person = parse_person_item(raw, True)
        if person:
            people.append(person)
    for raw in re.split(r"[、,，\n]+", substitutes_text or ""):
        person = parse_person_item(raw, False)
        if person:
            people.append(person)
    return people


def parse_events(text: str, event_type: str) -> list[dict]:
    if not text:
        return []

    events = []
    for line in re.split(r"\n+", text):
        line = clean_text(line)
        if not line:
            continue

        # Examples in the sheet include "1、5.许景粟", "1, 3.蒋思杨",
        # and "-5.许景粟". The numbers are scoring-order labels.
        match = re.match(r"^([\-0-9\s,，、]+)[.。]\s*(.+)$", line)
        if match:
            numbers = [int(num) for num in re.findall(r"\d+", match.group(1))]
            name = normalize_name(match.group(2))
        else:
            numbers = []
            name = normalize_name(line)

        if not name:
            continue
        if not numbers:
            numbers = [None]

        for number in numbers:
            events.append(
                {
                    "order": number,
                    "player": name,
                    "type": event_type,
                    "ownGoal": name.upper() in {"OG", "OWNGOAL", "OWNGOAL."},
                }
            )
    return events


def empty_player_stats() -> dict:
    return {
        "goals": 0,
        "assists": 0,
        "appearances": 0,
        "starts": 0,
        "wins": 0,
        "draws": 0,
        "losses": 0,
        "cleanSheets": 0,
        "goalsAgainst": 0,
        "knownResults": 0,
        "goalkeeperAppearances": 0,
    }


def empty_team_stats() -> dict:
    return {
        "matches": 0,
        "knownScoreMatches": 0,
        "goalsFor": 0,
        "goalsAgainst": 0,
        "wins": 0,
        "draws": 0,
        "losses": 0,
        "cleanSheets": 0,
    }


def ensure_player(stats: dict, name: str) -> dict:
    if name not in stats:
        stats[name] = empty_player_stats()
    return stats[name]


def add_match_to_stats(stats_bucket: dict, match: dict) -> None:
    team_stats = stats_bucket["team"]
    player_stats = stats_bucket["players"]

    score = match["score"]
    known_score = score["known"]
    if known_score:
        team_stats["matches"] += 1
        team_stats["knownScoreMatches"] += 1
        team_stats["goalsFor"] += score["goalsFor"]
        team_stats["goalsAgainst"] += score["goalsAgainst"]
        if score["goalsAgainst"] == 0:
            team_stats["cleanSheets"] += 1
        if score["result"] == "win":
            team_stats["wins"] += 1
        elif score["result"] == "draw":
            team_stats["draws"] += 1
        elif score["result"] == "loss":
            team_stats["losses"] += 1
    else:
        team_stats["matches"] += 1

    for person in match["roster"]:
        player = ensure_player(player_stats, person["name"])
        player["appearances"] += 1
        if person["starter"]:
            player["starts"] += 1
        if known_score:
            player["knownResults"] += 1
            if score["result"] == "win":
                player["wins"] += 1
            elif score["result"] == "draw":
                player["draws"] += 1
            elif score["result"] == "loss":
                player["losses"] += 1

            if person["goalkeeper"]:
                player["goalkeeperAppearances"] += 1
                player["goalsAgainst"] += score["goalsAgainst"]
                if score["goalsAgainst"] == 0:
                    player["cleanSheets"] += 1

    for event in match["goals"]:
        if event["ownGoal"]:
            continue
        player = ensure_player(player_stats, event["player"])
        player["goals"] += 1

    for event in match["assists"]:
        if event["ownGoal"]:
            continue
        player = ensure_player(player_stats, event["player"])
        player["assists"] += 1


def finalize_player_stats(raw_players: dict) -> list[dict]:
    players = []
    for name, stats in raw_players.items():
        known_results = stats["knownResults"]
        win_rate = stats["wins"] / known_results if known_results else None
        players.append({"name": name, **stats, "winRate": win_rate})
    return sorted(
        players,
        key=lambda p: (
            -p["goals"],
            -p["assists"],
            -p["appearances"],
            p["name"],
        ),
    )


def build_stats(matches: list[dict], seasons: list[str]) -> dict:
    buckets = {
        "all": {"team": empty_team_stats(), "players": {}},
        "bySeason": {season: {"team": empty_team_stats(), "players": {}} for season in seasons},
    }

    for match in matches:
        add_match_to_stats(buckets["all"], match)
        add_match_to_stats(buckets["bySeason"][match["season"]], match)

    return {
        "all": {
            "team": buckets["all"]["team"],
            "players": finalize_player_stats(buckets["all"]["players"]),
        },
        "bySeason": {
            season: {
                "team": bucket["team"],
                "players": finalize_player_stats(bucket["players"]),
            }
            for season, bucket in buckets["bySeason"].items()
        },
    }


def meta_for_player(player_meta: dict, season: str, name: str) -> dict:
    return player_meta.get((season, name), {})


def build_rosters(matches: list[dict], seasons: list[str], player_meta: dict | None = None) -> dict:
    player_meta = player_meta or {}
    rosters: dict[str, dict[str, dict]] = {season: {} for season in seasons}
    for match in matches:
        season_roster = rosters[match["season"]]
        for person in match["roster"]:
            meta = meta_for_player(player_meta, match["season"], person["name"])
            player = season_roster.setdefault(
                person["name"],
                {
                    "name": person["name"],
                    "position": meta.get("position") or player_position(person["name"]),
                    "captain": bool(meta.get("captain")),
                    "number": meta.get("number") or player_number(match["season"], person["name"]),
                    "photo": meta.get("photo", ""),
                    "appearances": 0,
                },
            )
            player["appearances"] += 1
            player["captain"] = bool(player["captain"] or meta.get("captain"))
            if person["goalkeeper"]:
                player["position"] = "门将"
            elif not player["position"]:
                player["position"] = player_position(person["name"])

        for event in [*match["goals"], *match["assists"]]:
            if event["ownGoal"]:
                continue
            meta = meta_for_player(player_meta, match["season"], event["player"])
            season_roster.setdefault(
                event["player"],
                {
                    "name": event["player"],
                    "position": meta.get("position") or player_position(event["player"]),
                    "captain": bool(meta.get("captain")),
                    "number": meta.get("number") or player_number(match["season"], event["player"]),
                    "photo": meta.get("photo", ""),
                    "appearances": 0,
                },
            )

    finalized = {}
    for season, players in rosters.items():
        for meta_season, meta_name in [key for key in player_meta if key[0] == season]:
            meta = player_meta[(meta_season, meta_name)]
            players.setdefault(
                meta_name,
                {
                    "name": meta_name,
                    "position": meta.get("position") or player_position(meta_name),
                    "captain": bool(meta.get("captain")),
                    "number": meta.get("number") or player_number(season, meta_name),
                    "photo": meta.get("photo", ""),
                    "appearances": 0,
                },
            )
        captain_name = SEASON_CAPTAINS.get(season, "")
        if captain_name:
            meta = meta_for_player(player_meta, season, captain_name)
            captain = players.setdefault(
                captain_name,
                {
                    "name": captain_name,
                    "position": meta.get("position") or player_position(captain_name),
                    "captain": False,
                    "number": meta.get("number") or player_number(season, captain_name),
                    "photo": meta.get("photo", ""),
                    "appearances": 0,
                },
            )
            captain["captain"] = True
        finalized[season] = sorted(players.values(), key=lambda p: (-p["captain"], p["name"]))
    return finalized


def truthy(value: str) -> bool:
    return clean_text(value).lower() in {"1", "true", "yes", "y", "是", "队长"}


def sheet_records(workbook, sheet_name: str) -> list[dict]:
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    headers = [clean_text(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = {headers[index]: clean_text(value) for index, value in enumerate(row[: len(headers)])}
        if any(record.values()):
            records.append(record)
    return records


def make_match(
    *,
    row_index: int,
    match_id: str,
    season: str,
    competition: str,
    date_text: str,
    venue: str,
    stage: str,
    opponent: str,
    score_raw: str,
    goals_raw: str,
    assists_raw: str,
    starting_raw: str,
    substitutes_raw: str,
    photo: str = "",
    notes: str = "",
) -> dict:
    starters_count = STARTERS_BY_COMPETITION.get(competition, 0)
    score = parse_score(score_raw)
    goals = parse_events(goals_raw, "goal")
    assists = parse_events(assists_raw, "assist")
    roster = parse_roster_columns(starting_raw, substitutes_raw)
    goal_count_from_events = len(goals)
    goals_incomplete = (
        score["known"]
        and score["goalsFor"] is not None
        and score["goalsFor"] > goal_count_from_events
    )

    return {
        "id": match_id or slugify(str(row_index), season, competition, stage, opponent),
        "sourceRow": row_index,
        "season": season,
        "competition": competition,
        "format": FORMAT_BY_COMPETITION.get(competition, ""),
        "startersCount": starters_count,
        "date": date_text,
        "venue": full_venue_name(venue),
        "stage": stage,
        "opponent": opponent,
        "homeTeam": TEAM_CN,
        "score": score,
        "goals": goals,
        "assists": assists,
        "roster": roster,
        "photo": photo,
        "notes": notes,
        "raw": {
            "goals": goals_raw,
            "assists": assists_raw,
            "startingLineup": starting_raw,
            "substitutes": substitutes_raw,
        },
        "dataQuality": {
            "missingDate": not bool(date_text),
            "missingVenue": not bool(venue),
            "missingStage": not bool(stage),
            "missingOpponent": not bool(opponent),
            "missingScore": not bool(score_raw),
            "missingRoster": not bool(roster),
            "goalsIncomplete": goals_incomplete,
            "missingAssists": score["known"] and score["goalsFor"] and not bool(assists),
        },
    }


def load_database() -> dict:
    workbook = openpyxl.load_workbook(DATABASE_FILE, data_only=True)

    competition_records = sheet_records(workbook, "Competitions")
    starters = {
        record["competition"]: int(record["starters_count"])
        for record in competition_records
        if record.get("competition") and record.get("starters_count")
    }
    formats = {
        record["competition"]: record["format"]
        for record in competition_records
        if record.get("competition") and record.get("format")
    }
    if starters:
        STARTERS_BY_COMPETITION.update(starters)
    if formats:
        FORMAT_BY_COMPETITION.update(formats)

    for record in sheet_records(workbook, "Venues"):
        if record.get("short_name") and record.get("full_name"):
            VENUE_NAMES[record["short_name"]] = record["full_name"]

    player_meta = {}
    for record in sheet_records(workbook, "Players"):
        season = record.get("season", "")
        name = normalize_name(record.get("name", ""))
        if not season or not name:
            continue
        player_meta[(season, name)] = {
            "number": record.get("number", ""),
            "position": record.get("position", ""),
            "captain": truthy(record.get("captain", "")),
            "photo": record.get("photo", ""),
        }

    matches = []
    seasons = []
    competitions = []
    for index, record in enumerate(sheet_records(workbook, "Matches"), start=2):
        season = record.get("season", "")
        competition = record.get("competition", "")
        if not season or not competition:
            continue
        if season not in seasons:
            seasons.append(season)
        if competition not in competitions:
            competitions.append(competition)
        matches.append(
            make_match(
                row_index=int(record.get("source_row") or index),
                match_id=record.get("match_id", ""),
                season=season,
                competition=competition,
                date_text=record.get("date", ""),
                venue=record.get("venue", ""),
                stage=record.get("stage", ""),
                opponent=record.get("opponent", ""),
                score_raw=record.get("score", ""),
                goals_raw=record.get("goals_raw", ""),
                assists_raw=record.get("assists_raw", ""),
                starting_raw=record.get("starting_lineup", ""),
                substitutes_raw=record.get("substitutes", ""),
                photo=record.get("photo", ""),
                notes=record.get("notes", ""),
            )
        )

    achievements = [
        {
            "year": int(record["year"]) if str(record.get("year", "")).isdigit() else record.get("year", ""),
            "competition": record.get("competition", ""),
            "title": record.get("title", ""),
            "titleEn": record.get("title_en", ""),
            "image": record.get("image", ""),
        }
        for record in sheet_records(workbook, "Achievements")
    ]
    news = [
        {
            "season": record.get("season", ""),
            "type": record.get("type", ""),
            "typeEn": record.get("type_en", ""),
            "title": record.get("title", ""),
            "titleEn": record.get("title_en", ""),
            "date": record.get("date", ""),
            "url": record.get("url", ""),
            "cover": record.get("cover", ""),
            "description": record.get("description", ""),
            "descriptionEn": record.get("description_en", ""),
        }
        for record in sheet_records(workbook, "News")
    ]

    transfers_by_season: dict[str, dict] = {}
    for record in sheet_records(workbook, "Transfers"):
        season = record.get("season", "")
        if not season:
            continue
        bucket = transfers_by_season.setdefault(season, {"season": season, "newcomers": [], "departures": []})
        item = {
            "name": record.get("name", ""),
            "date": record.get("date", ""),
            "from": record.get("from", ""),
            "to": record.get("to", ""),
            "notes": record.get("notes", ""),
        }
        if record.get("direction", "").lower() == "out":
            bucket["departures"].append(item)
        else:
            bucket["newcomers"].append(item)

    return {
        "matches": matches,
        "seasons": seasons,
        "competitions": competitions,
        "achievements": achievements or ACHIEVEMENTS,
        "news": news or NEWS_ITEMS,
        "transfers": list(transfers_by_season.values()) or TRANSFERS,
        "playerMeta": player_meta,
    }


def main() -> None:
    player_meta = {}
    if DATABASE_FILE.exists():
        loaded = load_database()
        matches = loaded["matches"]
        seasons = loaded["seasons"]
        competitions = loaded["competitions"]
        achievements = loaded["achievements"]
        news = loaded["news"]
        transfers = loaded["transfers"]
        player_meta = loaded["playerMeta"]
        source_label = "whumath_database.xlsx"
    else:
        workbook = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
        sheet = workbook.worksheets[0]

        current_season = ""
        current_competition = ""
        matches = []
        seasons = []
        competitions = []
        achievements = ACHIEVEMENTS
        news = NEWS_ITEMS
        transfers = TRANSFERS
        source_label = "data.xlsx / 比赛数据"

        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            cells = [clean_text(cell) for cell in row[:11]]
            season_cell, competition_cell = cells[0], cells[1]
            if season_cell:
                current_season = season_cell
                if current_season not in seasons:
                    seasons.append(current_season)
            if competition_cell:
                current_competition = competition_cell
                if current_competition not in competitions:
                    competitions.append(current_competition)

            date_text, venue, stage, opponent, score_raw, goals_raw, assists_raw, starting_raw, substitutes_raw = cells[2:11]
            has_match_data = any([date_text, venue, stage, opponent, score_raw, goals_raw, assists_raw, starting_raw, substitutes_raw])
            if not has_match_data:
                continue

            matches.append(
                make_match(
                    row_index=row_index,
                    match_id=slugify(str(row_index), current_season, current_competition, stage, opponent),
                    season=current_season,
                    competition=current_competition,
                    date_text=date_text,
                    venue=venue,
                    stage=stage,
                    opponent=opponent,
                    score_raw=score_raw,
                    goals_raw=goals_raw,
                    assists_raw=assists_raw,
                    starting_raw=starting_raw,
                    substitutes_raw=substitutes_raw,
                )
            )

    data = {
        "meta": {
            "team": TEAM_CN,
            "teamEn": TEAM_EN,
            "source": source_label,
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "missingLabel": "待补充",
        },
        "seasons": seasons,
        "competitions": competitions,
        "formats": FORMAT_BY_COMPETITION,
        "startersByCompetition": STARTERS_BY_COMPETITION,
        "seasonCaptains": SEASON_CAPTAINS,
        "jerseyNumbers": JERSEY_NUMBERS,
        "achievements": achievements,
        "news": news,
        "transfers": transfers,
        "matches": matches,
        "stats": build_stats(matches, seasons),
        "rosters": build_rosters(matches, seasons, player_meta),
    }

    JSON_FILE.parent.mkdir(parents=True, exist_ok=True)
    JSON_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    JS_FILE.write_text(
        "window.WHUMATH_DATA = "
        + json.dumps(data, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )

    print(f"Wrote {JSON_FILE.relative_to(ROOT)}")
    print(f"Wrote {JS_FILE.relative_to(ROOT)}")
    print(f"Matches: {len(matches)}")


if __name__ == "__main__":
    main()
