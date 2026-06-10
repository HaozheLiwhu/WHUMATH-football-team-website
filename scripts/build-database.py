from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = ROOT / "data.xlsx"
DATABASE_FILE = ROOT / "whumath_database.xlsx"
IMAGE_DIR = ROOT / "assets" / "images"
NEWS_METADATA_FILE = ROOT / "data" / "news_metadata.json"

TEAM_CN = "数院"
TEAM_FULL_CN = "武大数院"

STARTERS_BY_COMPETITION = {
    "新生杯": 8,
    "振兴杯": 11,
    "五人制": 5,
}

FORMAT_BY_COMPETITION = {
    "新生杯": "8-a-side",
    "振兴杯": "11-a-side",
    "五人制": "5-a-side",
}

VENUE_NAMES = {
    "桂操": "桂园操场",
    "信操": "信息学部操场",
    "工操": "工学部操场",
    "奥场": "九一二操场",
}

SEASON_CAPTAINS = {
    "2023-2024": "胡然",
    "2024-2025": "廖梓瑞",
    "2025-2026": "李昊哲",
}

JERSEY_NUMBERS = {
    "2023-2024": {"李昊哲": "10"},
    "2024-2025": {"李昊哲": "10"},
    "2025-2026": {"李昊哲": "10"},
}

PLAYER_POSITIONS = {
    "李昊哲": "前腰",
}

ACHIEVEMENTS = [
    [2024, "五人制", "2024年五人制冠军", "2024 Five-a-side Champion", "assets/images/champion.jpg", ""],
    [2026, "五人制", "2026年五人制亚军", "2026 Five-a-side Runner-up", "", ""],
]

NEWS_URLS = [
    "https://mp.weixin.qq.com/s/OeadtfMaEzsPnL7Yyk_HDA",
    "https://mp.weixin.qq.com/s/gXa9S-cMVKNhL7aWZwCK9Q",
    "https://mp.weixin.qq.com/s/brT7mroZQHgZMFJ1ay2ZAA",
    "https://mp.weixin.qq.com/s/3m9qmbgsjCAzTsCAP1TSVA",
    "https://mp.weixin.qq.com/s/sHy6rgQglgoG6-g6Lk-WDA",
    "https://mp.weixin.qq.com/s/cNGJMCjo_PW1nEDck98B3Q",
    "https://mp.weixin.qq.com/s/3DRuuww1iuSTzWbb8Z5saQ",
    "https://mp.weixin.qq.com/s/hqF8HelH4fEthJbB8hSZDA",
]

def metadata_text(value) -> str:
    return str(value or "").strip()


def is_placeholder_text(value: str) -> bool:
    text = metadata_text(value)
    return (
        not text
        or bool(re.fullmatch(r"\?+\s*\d*", text))
        or text in {"微信公众平台", "环境异常", "安全验证"}
    )


def metadata_value(value, fallback: str = "") -> str:
    text = metadata_text(value)
    return fallback if is_placeholder_text(text) else text


def news_category(title: str, raw_type: str = "") -> tuple[str, str]:
    existing = metadata_text(raw_type)
    if existing in {"战报", "录像", "专访"}:
        category = existing
    elif "专访" in title:
        category = "专访"
    elif "录像" in title or "视频" in title:
        category = "录像"
    else:
        category = "战报"
    english = {"战报": "Match Report", "录像": "Video", "专访": "Interview"}[category]
    return category, english


def season_from_date(date_text: str, fallback: str = "2025-2026") -> str:
    match = re.match(r"^(20\d{2})\.(\d{2})\.\d{2}$", metadata_text(date_text))
    if not match:
        return fallback
    year = int(match.group(1))
    month = int(match.group(2))
    if month >= 9:
        return f"{year}-{year + 1}"
    if month <= 6:
        return f"{year - 1}-{year}"
    return f"{year} 暑假"


def date_sort_value(date_text: str) -> int:
    digits = re.sub(r"\D", "", metadata_text(date_text))
    return int(digits) if digits else 0


def local_news_cover(index: int) -> str:
    for extension in (".jpg", ".png", ".webp", ".gif"):
        path = IMAGE_DIR / f"news-{index:02d}{extension}"
        if path.exists():
            return path.relative_to(ROOT).as_posix()
    return ""


def existing_asset_path(value: str) -> str:
    text = metadata_text(value)
    if text and (ROOT / text).exists():
        return text.replace("\\", "/")
    return ""


def build_news_items() -> list[list]:
    metadata_by_url = {}
    if NEWS_METADATA_FILE.exists():
        metadata = json.loads(NEWS_METADATA_FILE.read_text(encoding="utf-8"))
        metadata_by_url = {item.get("url", ""): item for item in metadata if item.get("url")}

    items = []
    for index, url in enumerate(NEWS_URLS, start=1):
        item = metadata_by_url.get(url, {})
        title = metadata_value(item.get("title"), f"球队资讯 {index}")
        title_en = metadata_value(item.get("titleEn"), title)
        date_text = metadata_text(item.get("date"))
        category, category_en = news_category(title, item.get("type"))
        items.append(
            [
                season_from_date(date_text, metadata_value(item.get("season"), "2025-2026")),
                category,
                category_en,
                title,
                title_en,
                date_text,
                url,
                existing_asset_path(item.get("cover") or item.get("image")) or local_news_cover(index),
                metadata_value(item.get("description")),
                metadata_value(item.get("descriptionEn")),
            ]
        )
    return sorted(items, key=lambda row: date_sort_value(row[5]), reverse=True)


NEWS_ITEMS = build_news_items()

TRANSFERS = [
    ["2025-2026", "in", "贺宇洲", "青训", TEAM_FULL_CN, "2025.09", ""],
    ["2025-2026", "in", "王嘉宁", "青训", TEAM_FULL_CN, "2025.09", ""],
    ["2025-2026", "in", "应锦程", "吉林大学", TEAM_FULL_CN, "2025.09", ""],
    ["2024-2025", "in", "孔维轩", "青训", TEAM_FULL_CN, "2024.09", ""],
    ["2024-2025", "in", "姚宇辰", "青训", TEAM_FULL_CN, "2024.09", ""],
    ["2024-2025", "in", "刘岳然", "青训", TEAM_FULL_CN, "2024.09", ""],
    ["2024-2025", "in", "张硕铠", "青训", TEAM_FULL_CN, "2024.09", ""],
    ["2024-2025", "in", "姚健", "青训", TEAM_FULL_CN, "2024.09", ""],
    ["2024-2025", "in", "方烁俨", "武大经管", TEAM_FULL_CN, "2025.02", ""],
    ["2023-2024", "in", "李昊哲", "武大物院", TEAM_FULL_CN, "2024.02", ""],
    ["2023-2024", "in", "许景粟", "武大化院", TEAM_FULL_CN, "2023.09", ""],
    ["2023-2024", "in", "刘恺源", "青训", TEAM_FULL_CN, "2023.09", ""],
    ["2023-2024", "in", "韩志豪", "青训", TEAM_FULL_CN, "2023.09", ""],
    ["2023-2024", "in", "邝钧皓", "青训", TEAM_FULL_CN, "2023.09", ""],
    ["2024-2025", "out", "邝钧皓", TEAM_FULL_CN, "退役", "2025.09", ""],
    ["2023-2024", "out", "周宇杰", TEAM_FULL_CN, "澳门大学", "2024.06", ""],
    ["2023-2024", "out", "胡然", TEAM_FULL_CN, "哥伦比亚大学", "2024.06", ""],
    ["2023-2024", "out", "石德铭", TEAM_FULL_CN, "南洋理工大学", "2024.06", ""],
    ["2023-2024", "out", "格桑次仁", TEAM_FULL_CN, "退役", "2024.06", ""],
    ["2023-2024", "out", "遵追桑布", TEAM_FULL_CN, "退役", "2024.06", ""],
    ["2025-2026", "out", "蒋思杨", TEAM_FULL_CN, "退役", "2026.06", ""],
    ["2025-2026", "out", "廖梓瑞", TEAM_FULL_CN, "新加坡国立大学", "2026.06", ""],
    ["2025-2026", "out", "徐锋庆", TEAM_FULL_CN, "退役", "2026.06", ""],
    ["2025-2026", "out", "张梦洋", TEAM_FULL_CN, "退役", "2026.06", ""],
]


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y.%m.%d")
    return str(value).strip()


def normalize_name(name: str) -> str:
    return re.sub(r"\s+", "", name.strip())


def split_people(text: str) -> list[str]:
    people = []
    for raw in re.split(r"[、,，\n]+", clean_text(text)):
        item = clean_text(raw)
        if not item:
            continue
        name = re.sub(r"[（(]\s*(C|GK|P)\s*[）)]", "", item, flags=re.IGNORECASE)
        name = normalize_name(name)
        if name:
            people.append(name)
    return people


def slugify(*parts: str) -> str:
    text = "-".join(part for part in parts if part)
    text = re.sub(r"\s+", "-", text.strip())
    text = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", text)
    return text.strip("-").lower()


def add_sheet(wb, name: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="0B2E72")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = max(len(str(ws.cell(row=row, column=col).value or "")) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
    ws.freeze_panes = "A2"


def build_match_rows() -> tuple[list[list], dict[str, set[str]]]:
    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    ws = wb.worksheets[0]
    current_season = ""
    current_competition = ""
    rows = []
    players_by_season: dict[str, set[str]] = {}

    for source_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cells = [clean_text(cell) for cell in row[:11]]
        season_cell, competition_cell = cells[0], cells[1]
        if season_cell:
            current_season = season_cell
            players_by_season.setdefault(current_season, set())
        if competition_cell:
            current_competition = competition_cell

        date_text, venue, stage, opponent, score, goals, assists, starters, substitutes = cells[2:11]
        has_match_data = any([date_text, venue, stage, opponent, score, goals, assists, starters, substitutes])
        if not has_match_data:
            continue

        venue = VENUE_NAMES.get(venue, venue)
        match_id = slugify(str(source_row), current_season, current_competition, stage, opponent)
        for name in split_people(starters) + split_people(substitutes):
            players_by_season.setdefault(current_season, set()).add(name)

        rows.append(
            [
                match_id,
                source_row,
                current_season,
                current_competition,
                date_text,
                venue,
                stage,
                opponent,
                score,
                goals,
                assists,
                starters,
                substitutes,
                "",
                "",
            ]
        )
    return rows, players_by_season


def build_player_rows(players_by_season: dict[str, set[str]]) -> list[list]:
    for season, direction, name, *_rest in TRANSFERS:
        players_by_season.setdefault(season, set()).add(name)

    rows = []
    for season in sorted(players_by_season):
        captain_name = SEASON_CAPTAINS.get(season, "")
        if captain_name:
            players_by_season[season].add(captain_name)
        for name in sorted(players_by_season[season]):
            rows.append(
                [
                    season,
                    name,
                    JERSEY_NUMBERS.get(season, {}).get(name, ""),
                    PLAYER_POSITIONS.get(name, ""),
                    "TRUE" if name == captain_name else "",
                    "",
                    "",
                ]
            )
    return rows


def main() -> None:
    match_rows, players_by_season = build_match_rows()
    player_rows = build_player_rows(players_by_season)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    add_sheet(
        wb,
        "Matches",
        [
            "match_id",
            "source_row",
            "season",
            "competition",
            "date",
            "venue",
            "stage",
            "opponent",
            "score",
            "goals_raw",
            "assists_raw",
            "starting_lineup",
            "substitutes",
            "photo",
            "notes",
        ],
        match_rows,
    )
    add_sheet(wb, "Players", ["season", "name", "number", "position", "captain", "photo", "notes"], player_rows)
    add_sheet(wb, "Transfers", ["season", "direction", "name", "from", "to", "date", "notes"], TRANSFERS)
    add_sheet(wb, "Achievements", ["year", "competition", "title", "title_en", "image", "notes"], ACHIEVEMENTS)
    add_sheet(
        wb,
        "News",
        ["season", "type", "type_en", "title", "title_en", "date", "url", "cover", "description", "description_en"],
        NEWS_ITEMS,
    )
    add_sheet(
        wb,
        "Competitions",
        ["competition", "format", "starters_count"],
        [[name, FORMAT_BY_COMPETITION[name], STARTERS_BY_COMPETITION[name]] for name in STARTERS_BY_COMPETITION],
    )
    add_sheet(wb, "Venues", ["short_name", "full_name"], [[k, v] for k, v in VENUE_NAMES.items()])
    add_sheet(
        wb,
        "README",
        ["sheet", "purpose"],
        [
            ["Matches", "每场比赛一行，首发和替补分列，GK/C 仍可写在姓名后括号中。"],
            ["Players", "赛季球员信息：号码、位置、队长、头像。"],
            ["Transfers", "转会记录：direction 为 in/out，from/to/date 控制页面流向。"],
            ["Achievements", "历史战绩与荣誉。"],
            ["News", "资讯、采访和外部链接。"],
            ["Competitions", "赛事赛制和首发人数。"],
            ["Venues", "场地简称到全名的映射。"],
        ],
    )

    wb.save(DATABASE_FILE)
    print(f"Wrote {DATABASE_FILE.relative_to(ROOT)}")
    print(f"Matches: {len(match_rows)}")
    print(f"Player season rows: {len(player_rows)}")


if __name__ == "__main__":
    main()
